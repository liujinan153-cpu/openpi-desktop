# -*- coding: utf-8 -*-
# 8.2 内容 + 0035 模板:基于 template0035.pptx,克隆补页、替换文本、重排页序
import copy
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.text.text import _Paragraph

SRC, OUT = 'template0035.pptx', 'build_out.pptx'
prs = Presentation(SRC)
slides = list(prs.slides)
assert len(slides) == 24

# ---------- helpers ----------
def shapes_by_id(slide):
    d = {}
    def walk(shapes):
        for s in shapes:
            d[s.shape_id] = s
            if s.shape_type == 6:
                walk(s.shapes)
    walk(slide.shapes)
    return d

def replace_in_paragraph(p, new_text):
    runs = p.runs
    if not runs:
        p.add_run().text = new_text
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)

def set_lines(tf, lines):
    paras = list(tf.paragraphs)
    donor = copy.deepcopy(paras[0]._p)
    txBody = tf._txBody
    for p in paras[1:]:
        txBody.remove(p._p)
    replace_in_paragraph(tf.paragraphs[0], lines[0])
    for line in lines[1:]:
        newp = copy.deepcopy(donor)
        txBody.append(newp)
        replace_in_paragraph(_Paragraph(newp, tf), line)

def set_text(slide, sid, lines):
    s = shapes_by_id(slide)[sid]
    set_lines(s.text_frame, lines if isinstance(lines, list) else [lines])

def clone_slide(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    spTree = new.shapes._spTree
    for el in list(spTree):
        if el.tag.split('}')[1] in ('sp', 'pic', 'graphicFrame', 'grpSp', 'cxnSp'):
            spTree.remove(el)
    for el in src.shapes._spTree:
        if el.tag.split('}')[1] in ('sp', 'pic', 'graphicFrame', 'grpSp', 'cxnSp'):
            spTree.append(copy.deepcopy(el))
    for el in spTree.iter():
        for attr in (qn('r:embed'), qn('r:link'), qn('r:id')):
            rid = el.get(attr)
            if rid:
                rel = src.part.rels[rid]
                if rel.is_external:
                    el.set(attr, new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref))
                else:
                    el.set(attr, new.part.relate_to(rel.target_part, rel.reltype))
    return new

IND = '      '
# ---------- 1. 先克隆( pristine 母版页) ----------
C1 = clone_slide(prs, slides[13])   # 做电子产品的主人
C2 = clone_slide(prs, slides[18])   # 选择题
C3 = clone_slide(prs, slides[19])   # 情境题1
C4 = clone_slide(prs, slides[19])   # 情境题2
C5 = clone_slide(prs, slides[21])   # 板书设计
C6 = clone_slide(prs, slides[23])   # 谢谢聆听

# ---------- 2. 文本替换(按页对象+shape_id) ----------
S = slides
# 封面
set_text(S[0], 2, '第八课《别让它抢走太多》第2课时')
set_text(S[0], 6, ['部编版道德与法治四年级上册', '第三单元《生活在信息社会》'])
# 情境导入 ← 8.2 S1
set_text(S[2], 10, '观察生活：哥哥错过了什么？')
set_text(S[2], 9, [IND + '弟弟抱着篮球热情邀请：“哥哥，下楼一起打篮球吧！”　哥哥紧盯着电视屏：“别打扰我，我还要看动画片呢！”',
                   '思考：沉迷电子屏幕，会让哥哥失去哪些宝贵的时光？'])
# 被抢走的第1份礼物 ← 8.2 S2
set_text(S[5], 10, '被抢走的第1份礼物：户外运动与同伴欢笑')
set_text(S[5], 9, [IND + '错过的快乐：缺席同伴游戏，失去锻炼身体的好机会。',
                   '健康的代价：缺乏运动导致体质下降，性格变得孤僻懒散。',
                   '视觉隐喻：沙漏中急速流逝的金色时光粒子，象征被屏幕吞噬的户外时光。'])
# 家庭危机 ← 8.2 S4
set_text(S[6], 10, '家庭危机：冷清的客厅与疏离的心')
set_text(S[6], 9, [IND + '现象诊断：爸爸刷新闻、妈妈逛电商、孩子玩游戏，一家人零交流。',
                   '情感损失：缺失温馨的谈心时刻，家庭氛围变得冷漠冰凉。',
                   '核心反思：屏幕连接了远方，却拉远了最近的亲人！'])
# 健康警报 ← 8.2 S5
set_text(S[7], 10, '健康警报：手机对身心的隐形伤害')
set_text(S[7], 9, [IND + '视力危机：长时间盯着屏幕，导致眼睛干涩、近视加深。',
                   '睡眠障碍：蓝光抑制褪黑素，引发失眠与白天疲倦。',
                   '自控力下降：大脑习惯高刺激，对现实学习生活失去兴趣。'])
# 活动1:四类时间 ← 8.2 S7(标题保留模板结构)
set_text(S[10], 9, [IND + '学习与写作业的时间：注意力不集中，拖延作业。',
                    '与父母交流的时间：各刷各的屏幕，家庭沟通变少。',
                    '充足睡眠的时间：深夜熬夜刷设备，白天上课瞌睡。',
                    '阅读课外书的时间：习惯快餐短视频，无法静心阅读。',
                    '句式：我不愿意被抢走______，因为______。'])
# 共读豆豆一家 ← 8.2 S9
set_text(S[11], 9, [IND + '做法借鉴：放下手机电视，每周设立固定“无电子设备家庭夜”。',
                    '活动形式：举办家庭故事会、棋牌游戏、户外散步谈心。',
                    '改变成果：拉近亲子距离，收获满满的欢笑与家庭温暖。'])
# 克隆页:做电子产品的主人 ← 8.2 S10
set_text(C1, 10, '做电子产品的主人，不做奴隶')
set_text(C1, 9, [IND + '工具属性：电子产品是学习辅助与适度娱乐的工具。',
                 '主动选择：主动留出线下时间，拥抱现实生活的精彩。',
                 '视觉隐喻：平衡天秤一端是小巧的手机，另一端是沉甸甸的爱心、书本与运动器材。'])
# 如何制定使用约定 ← 8.2 S11
set_text(S[15], 10, '如何制定“使用约定”')
set_text(S[15], 9, [IND + '限定使用时长：每天/每周末累计使用时间（网课+娱乐不超限）。',
                    '明确使用场景：吃饭、睡前、全家聚餐时坚决不碰设备。',
                    '严格筛选内容：优先看益智科普、健康动画，拒绝不良信息。'])
# 起草约定 ← 8.2 S12
set_text(S[16], 9, [IND + '同桌合作：讨论并拟定3条切实可行的家庭电子产品使用规则。',
                    '互查互评：条款是否具体？执行是否困难？能否互相监督？',
                    '教师点评：避免“完全禁止”或“过于宽泛”，确保契约可落地。'])
# 克隆页:选择题 ← 8.2 S14
set_text(C2, 10, '课堂练习：选择题')
set_text(C2, 9, [IND + '1. 好朋友约你下楼跳绳，你正好看动画片，正确的做法是（　　）',
                 'A. 看完这一集再下楼　　B. 关掉电视，立刻出门和朋友玩耍　　C. 拒绝朋友，继续在家看电视',
                 '2. 每天使用平板上网课、看动画，累计时长最好不超过（　　）',
                 'A. 2小时　　B. 半小时　　C. 5小时',
                 '3. 晚饭全家聚餐时，我们应该（　　）',
                 'A. 边吃饭边刷短视频　　B. 放下手机，和家人聊天　　C. 抱着平板边吃边看',
                 '参考答案：B、A、B'])
# 克隆页:情境题1 ← 8.2 S16
set_text(C3, 10, '课堂练习：情境题1')
set_text(C3, 9, [IND + '情境：晚上妈妈想和你聊聊学校的趣事，你正在刷短视频舍不得放下手机。你会怎么做？',
                 '参考答案：立刻放下手机，认真听妈妈聊天，主动和妈妈分享学校发生的事，约定聊天结束后再短暂使用电子产品。'])
# 克隆页:情境题2 ← 8.2 S17
set_text(C4, 10, '课堂练习：情境题2')
set_text(C4, 9, [IND + '情境：周末你写完作业，想长时间看动画片，爸爸提醒你要出门打球。你怎么安排时间？',
                 '参考答案：和爸爸约定看20分钟动画片，看完后放下电视，出门和爸爸打球运动，平衡娱乐和运动时间。'])
# 课堂总结 ← 8.2 S13
set_text(S[21], 9, [IND + '同学们，今天我们懂得电子产品虽有趣，却不能挤占我们运动、陪伴家人、学习休息的时间。豆豆一家的故事告诉我们，放下电子设备才能感受家庭温暖。',
                    '课后大家要和父母一起制定电子产品使用约定，学会自律管控，多走进现实生活，珍惜和同伴、家人相处的美好时光。'])
# 克隆页:板书设计 ← 8.2 S18
set_text(C5, 10, '板书设计')
set_text(C5, 9, [IND + '8.2 我和它有个“约定”',
                 '1. 电子产品会抢走什么？——玩耍时间｜亲子交流时间｜学习、休息、运动时间',
                 '2. 榜样：豆豆家庭故事会——放下电子设备，珍惜线下陪伴',
                 '3. 我的约定：自律控时长、选健康内容、多线下生活'])
# 课后作业 ← 8.2 S19
set_text(S[23], 9, [IND + '实践任务：带上课堂设计的《约定单》，与父母召开家庭会议。',
                    '共同签署：完善细节，全员签字，贴在客厅显眼位置。',
                    '习惯养成：坚持执行一周，记录自己的自控心得与家庭变化。',
                    '拓展常识：坚持“20-20-20”护眼法则，做健康有节制的信息时代好少年！'])
# 克隆页:谢谢聆听 ← 8.2 S20
set_text(C6, 10, '谢谢聆听·下课！')
set_text(C6, 9, [IND + '拒绝沉迷 · 自律生活',
                 '做电子产品的主人，拥抱无比精彩的现实世界！'])

# ---------- 3. 重排页序 ----------
order = [S[0], S[1], S[2], S[3], S[4], S[5], S[6], S[7], S[8], S[9], S[10], S[11], S[12], S[13],
         C1,
         S[14], S[15], S[16], S[17], S[18], S[19], S[20],
         C2, C3, C4,
         S[21], C5,
         S[22],
         S[23], C6]
sldIdLst = prs.slides._sldIdLst
part2elem = {}
for sldId in list(sldIdLst):
    rId = sldId.get(qn('r:id'))
    part2elem[prs.part.rels[rId].target_part] = sldId
for s in order:
    sldIdLst.append(part2elem[s.part])  # append 已有元素 = 移动到末尾,最终即目标顺序

prs.save(OUT)
print('saved', OUT, 'slides =', len(Presentation(OUT).slides))
