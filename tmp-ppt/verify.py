# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.oxml.ns import qn

prs = Presentation('build_out.pptx')
print('slides =', len(prs.slides))
problems = []
for i, slide in enumerate(prs.slides):
    texts = []
    def walk(shapes):
        for s in shapes:
            if s.shape_type == 6:
                walk(s.shapes)
                continue
            # 图片 rId 解析检查
            for blip in s._element.iter(qn('a:blip')):
                rid = blip.get(qn('r:embed'))
                if rid and rid not in slide.part.rels:
                    problems.append(f'slide{i} broken image rel {rid}')
            if s.has_text_frame and s.text_frame.text.strip():
                texts.append(s.text_frame.text.replace('\n', ' / ').replace('\x0b', ' '))
    walk(slide.shapes)
    first = texts[0] if texts else '(无文本)'
    print(f'--- {i}: {first[:60]}')
    for t in texts[1:3]:
        print(f'      {t[:70]}')

# 占位残留检查
import zipfile, re
z = zipfile.ZipFile('build_out.pptx')
for n in z.namelist():
    if n.startswith('ppt/slides/slide') and n.endswith('.xml'):
        x = z.read(n).decode('utf8')
        for kw in ['Click to add', 'lorem', 'TODO', '[insert']:
            if kw in x:
                problems.append(f'{n}: leftover {kw}')
print('\nPROBLEMS:', problems if problems else 'NONE')
