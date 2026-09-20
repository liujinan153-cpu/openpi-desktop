import sys
from pptx import Presentation
from pptx.util import Emu

def iter_shapes(shapes, depth=0):
    for s in shapes:
        yield s, depth
        if s.shape_type == 6:
            yield from iter_shapes(s.shapes, depth+1)

def cm(v): return round(Emu(v).cm,1) if v is not None else None

path = sys.argv[1]
prs = Presentation(path)
print(f"SLIDE SIZE: {prs.slide_width}x{prs.slide_height} ({cm(prs.slide_width)}x{cm(prs.slide_height)} cm) slides={len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    print(f"\n===== SLIDE {i} layout={slide.slide_layout.name} =====")
    for s, d in iter_shapes(slide.shapes):
        ind = "  "*d
        t = s.shape_type
        info = f"{ind}[{s.shape_id}] {t} name={s.name!r} pos=({cm(s.left)},{cm(s.top)}) size=({cm(s.width)}x{cm(s.height)})"
        if s.shape_type == 13:  # picture
            try: info += f" IMG={s.image.filename or s.image.content_type} {s.image.size}"
            except: pass
        print(info)
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                if p.text.strip():
                    sizes = [r.font.size.pt if r.font.size else None for r in p.runs]
                    fonts = list({r.font.name for r in p.runs if r.font.name})
                    print(f"{ind}   TXT: {p.text!r} sizes={sizes} fonts={fonts}")
